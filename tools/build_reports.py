from copy import deepcopy
from pathlib import Path
import csv
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
TEMPLATE = REPORTS / "course_report_template.docx"


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def count_files(path):
    return sum(1 for p in path.iterdir() if p.is_file())


def fmt(v):
    try:
        return f"{float(v):.4f}"
    except (ValueError, TypeError):
        return str(v)


def set_run(run, size=12, bold=False, color="000000"):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(paragraph, text, size=12, bold=False, color="000000", align=None):
    paragraph.text = ""
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return paragraph


def fill_cell(cell, heading, body):
    cell.text = ""
    p = cell.paragraphs[0]
    set_para(p, heading, size=14, bold=True)
    for line in body:
        p = cell.add_paragraph()
        set_para(p, line, size=10.5)


def table_lines(rows, headers):
    widths = [max(len(str(row.get(h, ""))) for row in rows + [dict(zip(headers, headers))]) for h in headers]
    out = [" | ".join(h.ljust(widths[i]) for i, h in enumerate(headers))]
    out.append("-+-".join("-" * w for w in widths))
    for row in rows:
        out.append(" | ".join(str(row.get(h, "")).ljust(widths[i]) for i, h in enumerate(headers)))
    return out


def clone_template(out, exp_label):
    shutil.copyfile(TEMPLATE, out)
    doc = Document(out)
    for p in doc.paragraphs:
        text = p.text
        if "实验XX" in text:
            set_para(p, exp_label, size=22, bold=True, color="FF0000", align=WD_ALIGN_PARAGRAPH.CENTER)
        elif "2026年" in text and "XX" in text:
            set_para(p, "2026年  9  月  15  日", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    return doc


def build_diabetes():
    base = ROOT / "diabetes_genetic_risk_monitoring"
    pc = base / "outputs" / "pc"
    lab = base / "outputs" / "lab_box"
    pc_metrics = read_csv(pc / "metrics.csv")
    lab_metrics = read_csv(lab / "metrics.csv")
    dist = read_csv(pc / "label_distribution.csv")

    out = REPORTS / "diabetes_genetic_risk_monitoring_report.docx"
    doc = clone_template(out, "实验一")
    t0, t1, t2 = doc.tables
    fill_cell(t0.cell(0, 0), "一、实验题目：", ["糖尿病遗传风险监测"])
    fill_cell(t0.cell(1, 0), "二、实验目的：", [
        "读取糖尿病遗传风险训练集与测试集，完成数据质量检查、缺失值处理和特征预处理。",
        "训练随机森林、LightGBM、CatBoost 三类模型，并使用软投票融合输出糖尿病风险概率。",
        "使用 Accuracy、Precision、Recall、F1、AUC 等指标评价模型效果，并比较 PC 端与实验箱端运行结果。"])
    fill_cell(t0.cell(2, 0), "三、实验内容：", [
        "1. 数据检查：训练集 5070 条、测试集 1000 条；训练集标签 0 为 3134 条，标签 1 为 1936 条。",
        "2. 缺失值处理：舒张压训练集缺失 247 条、测试集缺失 49 条，使用训练集拟合的预处理流程填补，避免数据泄漏。",
        "3. 模型训练：分别训练随机森林、LightGBM、CatBoost，并融合三个模型的预测概率。",
        "4. 结果输出：保存验证指标、测试集预测概率、ROC 曲线和混淆矩阵。"])
    fill_cell(t0.cell(3, 0), "四、算法原理", [
        "本实验属于典型的结构化医学数据二分类任务。输入特征包括性别、出生年份、体重指数、糖尿病家族史、舒张压、口服耐糖量测试、胰岛素释放实验、肱三头肌皮褶厚度等，输出为是否患有糖尿病。由于不同字段量纲不同且存在少量缺失值，训练前需要先完成缺失值填补和必要的编码处理，并且预处理参数只能由训练集拟合得到，再应用到验证集和测试集。",
        "随机森林是 Bagging 思想下的树模型集成方法。它通过有放回抽样构造多棵决策树，并在节点划分时随机选择部分特征，从而降低单棵决策树对样本扰动和噪声的敏感性。对于本实验这类特征数量不多、特征含义较明确的表格数据，随机森林可以作为稳定基线，帮助判断数据本身是否具备可分性。",
        "LightGBM 和 CatBoost 都属于梯度提升树模型。梯度提升的核心思想是后一棵树拟合前面模型尚未解释好的残差，逐步降低损失函数。LightGBM 训练效率较高，适合快速获得强基线；CatBoost 对类别特征和非线性组合处理较友好，可以补充 LightGBM 和随机森林在样本划分上的差异。",
        "软投票融合不直接选择某一个模型，而是把多个模型输出的糖尿病风险概率进行平均，再根据阈值得到最终类别。这样做的好处是降低单一模型偶然偏差，使最终预测更稳定。本实验同时使用 Accuracy、Precision、Recall、F1 和 AUC 评价模型，其中 Recall 更关注阳性样本检出能力，AUC 更关注风险概率排序能力。"])
    rows = []
    for env, metrics in [("PC", pc_metrics), ("实验箱", lab_metrics)]:
        for r in metrics:
            rows.append({"环境": env, "模型": r["模型"], "Acc": fmt(r["Accuracy"]), "Prec": fmt(r["Precision"]), "Recall": fmt(r["Recall"]), "F1": fmt(r["F1"]), "AUC": fmt(r["AUC"])})
    fill_cell(t1.cell(0, 0), "五、实验步骤：", [
        "1. 将 train.csv 和 test.csv 放入 diabetes_genetic_risk_monitoring/data。",
        "2. 运行 run_diabetes_risk.py，完成数据读取、缺失值检查、预处理、训练验证划分和模型训练。",
        "3. 分别评估随机森林、LightGBM、CatBoost 和软投票融合模型。",
        "4. 在 PC 端和实验箱端各运行一次，整理输出文件到 outputs/pc 与 outputs/lab_box。",
        "",
        "主要验证指标如下：",
        *table_lines(rows, ["环境", "模型", "Acc", "Prec", "Recall", "F1", "AUC"]),
        "",
        "实验箱端软投票融合结果最好：Accuracy=0.9606，Recall=0.9612，F1=0.9490，AUC=0.9920。"])
    fill_cell(t2.cell(1, 0), "问题解答：", [
        "1. 为什么要避免数据泄漏：预处理参数必须只在训练集上拟合，再应用到验证集和测试集，否则验证指标会虚高。",
        "2. 为什么使用多模型比较：不同树模型对特征组合和噪声的敏感性不同，比较后可选择更稳定的方案。",
        "3. 为什么采用软投票：软投票利用概率信息融合多个模型，实验箱端 F1 达到 0.9490，综合表现优于单模型。"])
    fill_cell(t2.cell(2, 0), "实验小结：", [
        "本次实验不是单纯把代码跑通，而是把一个结构化医学数据任务从 Windows 端整理到实验箱端完整复现了一遍。最开始遇到的问题主要集中在数据和目录上：原始文件名、输出文件夹和实验箱路径不统一，容易出现脚本找不到 train.csv、test.csv 或者结果文件分散的问题。因此我们先把项目整理成 diabetes_genetic_risk_monitoring 这样的英文目录，把 data、outputs/pc、outputs/lab_box 分开，后面再通过 GitHub 同步，实验箱 clone 后结构就清楚很多。",
        "建模过程中我体会比较深的是，结构化医学数据不一定需要一开始就上很复杂的深度学习模型。这个任务的字段都是表格特征，样本量也不算特别大，随机森林、LightGBM、CatBoost 这类树模型更合适，训练速度快，可解释性也更容易讲清楚。随机森林提供稳定基线，LightGBM 的 AUC 最高，说明概率排序能力强；CatBoost 表现也很接近。最后采用软投票，是因为它综合了三种模型的概率输出，实验箱端 F1 达到 0.9490，Recall 也达到 0.9612，说明对糖尿病阳性样本的检出更充分。",
        "实验里还提醒我，数据泄漏是机器学习实验中很容易忽略但很关键的问题。舒张压在训练集和测试集中都有缺失，如果在全数据上一起拟合填补规则，验证结果可能会被人为抬高。因此预处理流程必须只用训练集拟合，再应用到其他数据。PC 端和实验箱端结果基本一致，也说明迁移包、依赖说明和输出整理是有效的。相比只追求一个高指标，这次实验让我更理解完整实验流程的重要性：数据检查、模型选择、环境迁移和结果复核都要能解释清楚，报告才真实可信。"])
    doc.save(out)
    return out


def build_chest():
    base = ROOT / "chest_xray_pneumonia_classification"
    pc = read_csv(base / "outputs" / "pc" / "metrics.csv")[0]
    lab = read_csv(base / "outputs" / "lab_box" / "metrics.csv")[0]
    splits = []
    for split in ["train", "val", "test"]:
        for label in ["NORMAL", "PNEUMONIA"]:
            splits.append(f"{split}/{label}: {count_files(base / 'data' / split / label)} 张")

    out = REPORTS / "chest_xray_pneumonia_classification_report.docx"
    doc = clone_template(out, "实验五")
    t0, t1, t2 = doc.tables
    fill_cell(t0.cell(0, 0), "一、实验题目：", ["胸部 X 光片肺炎分类"])
    fill_cell(t0.cell(1, 0), "二、实验目的：", [
        "读取胸部 X 光图像数据，完成 NORMAL 与 PNEUMONIA 二分类模型训练。",
        "掌握图像预处理、迁移学习、训练验证监控和测试集评估流程。",
        "比较 PC 端 CUDA 环境下 Vision Transformer 与实验箱 CPU 环境下 ConvNeXt Tiny 的运行差异。"])
    fill_cell(t0.cell(2, 0), "三、实验内容：", [
        "1. 数据组织：" + "；".join(splits) + "。",
        "2. 图像预处理：统一尺寸、张量化和归一化，构建训练集、验证集和测试集 DataLoader。",
        "3. 模型训练：PC 端使用 vit_b_16 训练 10 轮；实验箱端使用 convnext_tiny 训练 5 轮。",
        "4. 指标评估：输出 Accuracy、Precision、Recall、F1、训练曲线、混淆矩阵和测试集预测结果。"])
    fill_cell(t0.cell(3, 0), "四、算法原理", [
        "本实验是医学图像二分类任务，输入为胸部 X 光图像，输出为 NORMAL 或 PNEUMONIA。图像分类与表格分类不同，模型需要从像素中自动学习纹理、边缘、局部阴影、肺野区域变化等特征。由于课程数据量较小，如果从零开始训练深度网络，很容易过拟合，因此实验采用迁移学习思路：加载在大规模图像数据上预训练过的模型，只替换最后的分类头，再针对肺炎二分类进行微调。",
        "PC 端使用 Vision Transformer。ViT 的核心思想是把一张图像切分为固定大小的 patch，每个 patch 类似自然语言处理中的 token，再加入位置编码后送入 Transformer 编码器。自注意力机制会计算不同 patch 之间的关联，因此模型不仅能关注局部纹理，也能学习肺部区域之间的全局关系。ViT 的计算量相对较大，更依赖 GPU，所以本机 CUDA 环境适合使用它进行实验。",
        "实验箱端使用 ConvNeXt Tiny。ConvNeXt 是在传统卷积网络基础上吸收现代网络设计思想得到的模型，仍然保留卷积对局部结构建模的优势，同时训练和推理开销比 ViT 更可控。实验箱端主要使用 CPU，若强行使用 ViT，训练时间会明显拉长，甚至影响实验能否顺利完成，因此选择 convnext_tiny 更符合迁移环境。",
        "评价指标方面，Accuracy 反映整体分类正确率，Precision 关注预测为肺炎的样本中有多少是真的肺炎，Recall 关注真实肺炎样本有多少被检出，F1 综合 Precision 和 Recall。医学筛查任务通常不能只看 Accuracy，因为漏检肺炎样本会带来更高风险，所以报告中特别比较了 Recall 和 F1。"])
    fill_cell(t1.cell(0, 0), "五、实验步骤：", [
        "1. 将 chest_xray_pneumonia_classification/data 按 train、val、test 以及 NORMAL、PNEUMONIA 子目录组织。",
        "2. 在 PC 端 conda 环境中确认 torch 2.7.1+cu118、CUDA 可用，选择 vit_b_16、batch size=8、lr=0.0001、epochs=10 运行。",
        "3. 在实验箱端使用 CPU 运行 convnext_tiny、batch size=8、lr=0.0001、epochs=5，保存输出结果。",
        "4. 汇总两端测试集指标和训练曲线，比较模型结构与硬件环境差异。",
        "",
        "测试集指标：",
        f"PC / {pc['模型']} / device={pc['设备']} / runtime={pc['运行秒数']}s / Accuracy={fmt(pc['Accuracy'])} / Precision={fmt(pc['Precision'])} / Recall={fmt(pc['Recall'])} / F1={fmt(pc['F1'])}",
        f"实验箱 / {lab['模型']} / device={lab['设备']} / runtime={lab['运行秒数']}s / Accuracy={fmt(lab['Accuracy'])} / Precision={fmt(lab['Precision'])} / Recall={fmt(lab['Recall'])} / F1={fmt(lab['F1'])}",
        "",
        "PC 端 ViT 的 F1 为 0.9020，高于实验箱端 ConvNeXt Tiny 的 0.8515；PC 端运行 10 轮仅 33.6 秒，实验箱 CPU 运行 5 轮耗时 397.65 秒。"])
    fill_cell(t2.cell(1, 0), "问题解答：", [
        "1. 为什么 PC 端使用 Vision Transformer：本 PC 机具备 CUDA GPU，ViT 能利用自注意力建模全局图像关系，且训练速度允许跑更多轮。",
        "2. 为什么实验箱端没有继续使用 ViT：实验箱端主要是 CPU 环境，ViT 计算量较大，训练耗时和内存压力更高，因此选用 convnext_tiny 保证可运行。",
        "3. 如何看待 Recall：Recall 表示肺炎样本被检出的比例，PC 端 Recall=0.8679，高于实验箱端 0.8113，漏检风险相对更低。"])
    fill_cell(t2.cell(2, 0), "实验小结：", [
        "本次实验的真实过程比单纯训练一个图像分类模型要复杂一些。开始时我们先按教程完成实验箱迁移准备，把数据集、脚本、依赖和输出目录整理到 GitHub 仓库。后来在 PC 端检查 conda 环境时发现本机其实有 CUDA，可以使用 torch 2.7.1+cu118 和 RTX 4070 Laptop GPU，因此模型选择不必完全受实验箱限制。于是 PC 端改用 vit_b_16 训练 10 轮，实验箱端则继续使用更稳妥的 convnext_tiny 在 CPU 上跑通流程。",
        "模型选型上的体会比较明显。ViT 是更前沿的视觉模型，能够通过自注意力学习不同图像 patch 之间的全局关系，用在胸片任务上有一定优势；但它对算力要求高，如果放到实验箱 CPU 上训练，时间成本会很大。ConvNeXt Tiny 虽然不是 Transformer，但结构轻一些，保留了卷积网络对局部纹理和边缘的建模能力，更适合实验箱端验证可迁移性。最终 PC 端 ViT 的 Accuracy 为 0.9099、F1 为 0.9020，实验箱端 ConvNeXt Tiny 的 Accuracy 为 0.8649、F1 为 0.8515，结果也符合这种硬件和模型差异。",
        "实验中遇到的问题主要有三个。第一是环境判断不能凭感觉，之前误以为本机不能用 CUDA，后来检查现有 conda 环境才确认 GPU 可用；第二是实验箱没有方便的浏览器和 GitHub 登录环境，结果传回时需要通过文件夹整理、分支或手动拷贝解决；第三是数据量比较小，验证集只有 16 张图像，训练过程中验证 F1 波动很大，所以最终分析不能只看某一轮验证结果，而要结合测试集指标和混淆矩阵。通过这次实验，我更理解医学 AI 实验不是模型越新越好，而是要把数据规模、硬件条件、运行时间、迁移难度和评价指标一起考虑，选择能真实落地并能解释清楚的方案。"])
    doc.save(out)
    return out


def main():
    REPORTS.mkdir(exist_ok=True)
    for path in [build_diabetes(), build_chest()]:
        print(path)


if __name__ == "__main__":
    main()
